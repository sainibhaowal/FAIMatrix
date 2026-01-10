"""
FAIM Unified Document Extraction - Production Grade
=====================================================

Complete text extraction from all document formats for FAIM ingestion.
Supports: PDF (with OCR + tables), DOCX, XLSX, PPTX, TXT, MD, JSON,
          CSV, Code files, Legacy Office (.doc, .xls, .ppt), EPUB, ODF, RTF

This file provides two interfaces:
1. extract_text(file_bytes, filename) -> Tuple[str, int]
   Used by: ingest_service.py (main pipeline)

2. extract_text_from_path(file_path, filename) -> Dict[str, Any]
   Used by: batch_upload.py (advanced with FAIM metadata/chunks)

All extractors are optimized for FAIM's Fractal Antisymmetric Inheritance Memory.
"""

import io
import json
import logging
import subprocess
from pathlib import Path
from typing import Any, Dict, List, Tuple

logger = logging.getLogger(__name__)

# =============================================================================
# FILE TYPE CATEGORIES
# =============================================================================

TEXT_EXTENSIONS = {
    ".txt",
    ".md",
    ".markdown",
    ".html",
    ".htm",
    ".xml",
    ".json",
    ".yaml",
    ".yml",
    ".csv",
    ".log",
    ".ini",
    ".cfg",
    ".toml",
    ".py",
    ".js",
    ".jsx",
    ".ts",
    ".tsx",
    ".java",
    ".c",
    ".cpp",
    ".h",
    ".hpp",
    ".cs",
    ".go",
    ".rs",
    ".sql",
    ".sh",
    ".bash",
    ".tex",
    ".rb",
    ".php",
    ".swift",
    ".kt",
    ".scala",
    ".vue",
    ".svelte",
    ".css",
    ".scss",
}

BINARY_DOCUMENT_EXTENSIONS = {
    ".pdf",
    ".docx",
    ".doc",
    ".pptx",
    ".ppt",
    ".xlsx",
    ".xls",
    ".odt",
    ".ods",
    ".rtf",
    ".epub",
}

# Language mapping for code files
LANGUAGE_MAP = {
    ".py": "python",
    ".js": "javascript",
    ".ts": "typescript",
    ".jsx": "jsx",
    ".tsx": "tsx",
    ".go": "go",
    ".rs": "rust",
    ".java": "java",
    ".c": "c",
    ".cpp": "cpp",
    ".cs": "csharp",
    ".rb": "ruby",
    ".php": "php",
    ".swift": "swift",
    ".kt": "kotlin",
    ".scala": "scala",
    ".sh": "bash",
    ".sql": "sql",
    ".yaml": "yaml",
    ".yml": "yaml",
    ".toml": "toml",
    ".xml": "xml",
    ".html": "html",
    ".css": "css",
    ".scss": "scss",
    ".vue": "vue",
    ".svelte": "svelte",
}


# =============================================================================
# PRIMARY INTERFACE 1: BYTES-BASED (for ingest_service.py)
# =============================================================================


def extract_text(file_bytes: bytes, filename: str) -> Tuple[str, int]:
    """
    Extract text content from a file.

    Args:
        file_bytes: Raw file content
        filename: Original filename (used to detect type)

    Returns:
        Tuple of (extracted_text, token_estimate)
    """
    ext = _get_extension(filename)

    if ext in TEXT_EXTENSIONS:
        return _extract_text_bytes(file_bytes, ext)
    elif ext == ".pdf":
        return _extract_pdf_bytes(file_bytes)
    elif ext in {".docx", ".doc"}:
        return _extract_docx_bytes(file_bytes, ext)
    elif ext in {".xlsx", ".xls"}:
        return _extract_excel_bytes(file_bytes, ext)
    elif ext in {".pptx", ".ppt"}:
        return _extract_pptx_bytes(file_bytes, ext)
    elif ext == ".epub":
        return _extract_epub_bytes(file_bytes)
    elif ext in {".odt", ".ods"}:
        return _extract_odf_bytes(file_bytes)
    elif ext == ".rtf":
        return _extract_rtf_bytes(file_bytes)
    else:
        # Try as plain text
        return _extract_text_bytes(file_bytes, ext)


# =============================================================================
# PRIMARY INTERFACE 2: PATH-BASED (for batch_upload.py with full metadata)
# =============================================================================


def extract_text_from_path(file_path: str, filename: str) -> Dict[str, Any]:
    """
    Extract text content from a file path with full FAIM metadata.

    Args:
        file_path: Path to the file on disk
        filename: Original filename

    Returns:
        Dict with: content, metadata, chunks (FAIM-ready), file_type
    """
    ext = Path(filename).suffix.lower()

    extractors = {
        # Text formats
        ".txt": _extract_text_path,
        ".md": _extract_markdown_path,
        ".markdown": _extract_markdown_path,
        # Data formats
        ".json": _extract_json_path,
        ".csv": _extract_csv_path,
        # PDF (enhanced with tables + OCR fallback)
        ".pdf": _extract_pdf_path_with_ocr,
        # Office documents (modern)
        ".docx": _extract_docx_path,
        ".xlsx": _extract_xlsx_path,
        ".pptx": _extract_pptx_path,
        # Office documents (legacy)
        ".doc": _extract_doc_legacy_path,
        ".xls": _extract_xls_legacy_path,
        ".ppt": _extract_ppt_legacy_path,
        # Code files
        **{ext: _extract_code_path for ext in LANGUAGE_MAP.keys()},
    }

    extractor = extractors.get(ext, _extract_text_path)

    try:
        result = extractor(file_path, filename)
        result["faim_metadata"] = {
            "source_file": filename,
            "file_type": ext,
            "extraction_method": extractor.__name__,
            "chunk_count": len(result.get("chunks", [])),
        }
        return result
    except Exception as e:
        logger.error(f"Failed to extract content from {filename}: {e}")
        return {
            "content": f"[Extraction Error: {str(e)}]",
            "metadata": {"error": str(e), "filename": filename},
            "chunks": [],
            "file_type": ext,
        }


# =============================================================================
# UTILITY FUNCTIONS
# =============================================================================


def _get_extension(filename: str) -> str:
    """Get lowercase file extension."""
    if "." not in filename:
        return ""
    return "." + filename.rsplit(".", 1)[-1].lower()


def _estimate_tokens(text: str) -> int:
    """Estimate token count (~4 chars per token)."""
    if not text:
        return 0
    return max(1, len(text) // 4)


def is_supported_file(filename: str) -> bool:
    """Check if file type is supported."""
    ext = _get_extension(filename)
    return ext in TEXT_EXTENSIONS or ext in BINARY_DOCUMENT_EXTENSIONS


# =============================================================================
# BYTES-BASED EXTRACTORS (for ingest_service.py)
# =============================================================================


def _extract_text_bytes(file_bytes: bytes, ext: str) -> Tuple[str, int]:
    """Extract from plain text files (bytes input)."""
    try:
        for encoding in ["utf-8", "latin-1", "cp1252"]:
            try:
                text = file_bytes.decode(encoding)
                break
            except UnicodeDecodeError:
                continue
        else:
            text = file_bytes.decode("utf-8", errors="ignore")

        if ext == ".json":
            try:
                data = json.loads(text)
                text = json.dumps(data, indent=2, ensure_ascii=False)
            except json.JSONDecodeError:
                pass

        tokens = _estimate_tokens(text)
        logger.info(f"[Extractor] Text file: {len(text)} chars, {tokens} tokens")
        return text, tokens

    except Exception as e:
        logger.error(f"[Extractor] Text extraction failed: {e}")
        return "", 0


def _extract_pdf_bytes(file_bytes: bytes) -> Tuple[str, int]:
    """Extract text from PDF using PyMuPDF or pdfplumber (bytes input)."""
    # Try PyMuPDF first
    try:
        import fitz

        doc = fitz.open(stream=file_bytes, filetype="pdf")
        text_parts = []
        for page_num, page in enumerate(doc):
            page_text = page.get_text()
            if page_text.strip():
                text_parts.append(f"--- Page {page_num + 1} ---\n{page_text}")
        doc.close()
        text = "\n\n".join(text_parts)
        tokens = _estimate_tokens(text)
        logger.info(f"[Extractor] PDF (pymupdf): {len(text)} chars, {tokens} tokens")
        return text, tokens
    except ImportError:
        pass
    except Exception as e:
        logger.warning(f"[Extractor] PyMuPDF failed: {e}")

    # Try pdfplumber
    try:
        import pdfplumber

        text_parts = []
        with pdfplumber.open(io.BytesIO(file_bytes)) as pdf:
            for i, page in enumerate(pdf.pages):
                page_text = page.extract_text() or ""
                tables = page.extract_tables() or []
                for table in tables:
                    for row in table:
                        page_text += " | ".join(str(cell or "") for cell in row) + "\n"
                if page_text.strip():
                    text_parts.append(f"--- Page {i + 1} ---\n{page_text}")
        text = "\n\n".join(text_parts)
        tokens = _estimate_tokens(text)
        logger.info(f"[Extractor] PDF (pdfplumber): {len(text)} chars, {tokens} tokens")
        return text, tokens
    except ImportError:
        pass
    except Exception as e:
        logger.warning(f"[Extractor] pdfplumber failed: {e}")

    # Fallback to pypdf
    try:
        try:
            from pypdf import PdfReader
        except ImportError:
            from PyPDF2 import PdfReader
        reader = PdfReader(io.BytesIO(file_bytes))
        text_parts = []
        for i, page in enumerate(reader.pages):
            page_text = page.extract_text() or ""
            if page_text.strip():
                text_parts.append(f"--- Page {i + 1} ---\n{page_text}")
        text = "\n\n".join(text_parts)
        tokens = _estimate_tokens(text)
        return text, tokens
    except Exception as e:
        logger.error(f"[Extractor] PDF extraction failed: {e}")
        return "", 0


def _extract_docx_bytes(file_bytes: bytes, ext: str) -> Tuple[str, int]:
    """Extract text from Word documents (bytes input)."""
    if ext == ".doc":
        # Legacy .doc - try via textract
        try:
            import tempfile

            import textract

            with tempfile.NamedTemporaryFile(suffix=".doc", delete=False) as f:
                f.write(file_bytes)
                temp_path = f.name
            text = textract.process(temp_path).decode("utf-8", errors="ignore")
            Path(temp_path).unlink(missing_ok=True)
            return text, _estimate_tokens(text)
        except Exception as e:
            logger.warning(f"[Extractor] Legacy .doc failed: {e}")
            return "", 0

    try:
        from docx import Document

        doc = Document(io.BytesIO(file_bytes))
        text_parts = []
        for para in doc.paragraphs:
            if para.text.strip():
                text_parts.append(para.text)
        for table in doc.tables:
            for row in table.rows:
                row_text = " | ".join(cell.text.strip() for cell in row.cells if cell.text.strip())
                if row_text:
                    text_parts.append(row_text)
        text = "\n\n".join(text_parts)
        tokens = _estimate_tokens(text)
        logger.info(f"[Extractor] DOCX: {len(text)} chars, {tokens} tokens")
        return text, tokens
    except ImportError:
        logger.warning("[Extractor] python-docx not installed")
        return "", 0
    except Exception as e:
        logger.error(f"[Extractor] DOCX extraction failed: {e}")
        return "", 0


def _extract_excel_bytes(file_bytes: bytes, ext: str) -> Tuple[str, int]:
    """Extract text from Excel files (bytes input)."""
    if ext == ".xls":
        try:
            import tempfile

            import xlrd

            with tempfile.NamedTemporaryFile(suffix=".xls", delete=False) as f:
                f.write(file_bytes)
                temp_path = f.name
            wb = xlrd.open_workbook(temp_path)
            text_parts = []
            for sheet_name in wb.sheet_names():
                sheet = wb.sheet_by_name(sheet_name)
                text_parts.append(f"=== Sheet: {sheet_name} ===")
                for row_idx in range(sheet.nrows):
                    row = [str(sheet.cell_value(row_idx, col_idx)) for col_idx in range(sheet.ncols)]
                    text_parts.append(" | ".join(row))
            Path(temp_path).unlink(missing_ok=True)
            text = "\n".join(text_parts)
            return text, _estimate_tokens(text)
        except ImportError:
            logger.warning("[Extractor] xlrd not installed for .xls")
            return "", 0
        except Exception as e:
            logger.error(f"[Extractor] XLS extraction failed: {e}")
            return "", 0

    try:
        import openpyxl

        wb = openpyxl.load_workbook(io.BytesIO(file_bytes), read_only=True)
        text_parts = []
        for sheet_name in wb.sheetnames:
            sheet = wb[sheet_name]
            text_parts.append(f"=== Sheet: {sheet_name} ===")
            for row in sheet.iter_rows(values_only=True):
                row_values = [str(cell) for cell in row if cell is not None]
                if row_values:
                    text_parts.append(" | ".join(row_values))
        wb.close()
        text = "\n".join(text_parts)
        tokens = _estimate_tokens(text)
        logger.info(f"[Extractor] Excel: {len(text)} chars, {tokens} tokens")
        return text, tokens
    except ImportError:
        logger.warning("[Extractor] openpyxl not installed")
        return "", 0
    except Exception as e:
        logger.error(f"[Extractor] Excel extraction failed: {e}")
        return "", 0


def _extract_pptx_bytes(file_bytes: bytes, ext: str) -> Tuple[str, int]:
    """Extract text from PowerPoint files (bytes input)."""
    if ext == ".ppt":
        try:
            import tempfile

            import textract

            with tempfile.NamedTemporaryFile(suffix=".ppt", delete=False) as f:
                f.write(file_bytes)
                temp_path = f.name
            text = textract.process(temp_path).decode("utf-8", errors="ignore")
            Path(temp_path).unlink(missing_ok=True)
            return text, _estimate_tokens(text)
        except Exception as e:
            logger.warning(f"[Extractor] Legacy .ppt failed: {e}")
            return "", 0

    try:
        from pptx import Presentation

        prs = Presentation(io.BytesIO(file_bytes))
        text_parts = []
        for slide_num, slide in enumerate(prs.slides, 1):
            slide_text = [f"--- Slide {slide_num} ---"]
            for shape in slide.shapes:
                if hasattr(shape, "text") and shape.text.strip():
                    slide_text.append(shape.text)
            if len(slide_text) > 1:
                text_parts.append("\n".join(slide_text))
        text = "\n\n".join(text_parts)
        tokens = _estimate_tokens(text)
        logger.info(f"[Extractor] PPTX: {len(text)} chars, {tokens} tokens")
        return text, tokens
    except ImportError:
        logger.warning("[Extractor] python-pptx not installed")
        return "", 0
    except Exception as e:
        logger.error(f"[Extractor] PPTX extraction failed: {e}")
        return "", 0


def _extract_epub_bytes(file_bytes: bytes) -> Tuple[str, int]:
    """Extract text from EPUB files (bytes input)."""
    try:
        import ebooklib
        from bs4 import BeautifulSoup
        from ebooklib import epub

        book = epub.read_epub(io.BytesIO(file_bytes))
        text_parts = []
        for item in book.get_items():
            if item.get_type() == ebooklib.ITEM_DOCUMENT:
                soup = BeautifulSoup(item.get_content(), "html.parser")
                text = soup.get_text(separator="\n", strip=True)
                if text:
                    text_parts.append(text)
        text = "\n\n".join(text_parts)
        tokens = _estimate_tokens(text)
        logger.info(f"[Extractor] EPUB: {len(text)} chars, {tokens} tokens")
        return text, tokens
    except ImportError:
        logger.warning("[Extractor] ebooklib not installed")
        return "", 0
    except Exception as e:
        logger.error(f"[Extractor] EPUB extraction failed: {e}")
        return "", 0


def _extract_odf_bytes(file_bytes: bytes) -> Tuple[str, int]:
    """Extract text from OpenDocument files (bytes input)."""
    try:
        from odf import text as odf_text
        from odf.opendocument import load

        doc = load(io.BytesIO(file_bytes))
        text_parts = []
        for para in doc.getElementsByType(odf_text.P):
            para_text = "".join(t.data for t in para.childNodes if hasattr(t, "data"))
            if para_text.strip():
                text_parts.append(para_text)
        text = "\n\n".join(text_parts)
        tokens = _estimate_tokens(text)
        logger.info(f"[Extractor] ODF: {len(text)} chars, {tokens} tokens")
        return text, tokens
    except ImportError:
        logger.warning("[Extractor] odfpy not installed")
        return "", 0
    except Exception as e:
        logger.error(f"[Extractor] ODF extraction failed: {e}")
        return "", 0


def _extract_rtf_bytes(file_bytes: bytes) -> Tuple[str, int]:
    """Extract text from RTF files (bytes input)."""
    try:
        from striprtf.striprtf import rtf_to_text

        rtf_content = file_bytes.decode("utf-8", errors="ignore")
        text = rtf_to_text(rtf_content)
        tokens = _estimate_tokens(text)
        logger.info(f"[Extractor] RTF: {len(text)} chars, {tokens} tokens")
        return text, tokens
    except ImportError:
        logger.warning("[Extractor] striprtf not installed")
        return "", 0
    except Exception as e:
        logger.error(f"[Extractor] RTF extraction failed: {e}")
        return "", 0


# =============================================================================
# PATH-BASED EXTRACTORS (for batch_upload.py with full FAIM metadata)
# =============================================================================


def _extract_text_path(file_path: str, filename: str) -> Dict[str, Any]:
    """Extract plain text file from path."""
    with open(file_path, "r", encoding="utf-8", errors="ignore") as f:
        content = f.read()
    return {
        "content": content,
        "metadata": {"type": "text", "filename": filename, "char_count": len(content)},
        "chunks": _chunk_text(content, source=filename),
        "file_type": "txt",
    }


def _extract_markdown_path(file_path: str, filename: str) -> Dict[str, Any]:
    """Extract markdown file with section awareness."""
    with open(file_path, "r", encoding="utf-8", errors="ignore") as f:
        content = f.read()

    sections = []
    current_section = {"title": "Introduction", "level": 0, "content": ""}

    for line in content.split("\n"):
        if line.startswith("#"):
            if current_section["content"].strip():
                sections.append(current_section)
            level = len(line) - len(line.lstrip("#"))
            title = line.lstrip("#").strip()
            current_section = {"title": title, "level": level, "content": ""}
        else:
            current_section["content"] += line + "\n"

    if current_section["content"].strip():
        sections.append(current_section)

    chunks = []
    for i, section in enumerate(sections):
        if section["content"].strip():
            chunks.append(
                {
                    "content": section["content"].strip(),
                    "metadata": {
                        "section": section["title"],
                        "section_level": section["level"],
                        "section_index": i,
                        "inheritance": f"section:{section['title']}",
                    },
                }
            )

    return {
        "content": content,
        "metadata": {"type": "markdown", "filename": filename, "sections": len(sections)},
        "chunks": chunks if chunks else _chunk_text(content, source=filename),
        "file_type": "markdown",
    }


def _extract_json_path(file_path: str, filename: str) -> Dict[str, Any]:
    """Extract JSON file with structure preservation."""
    with open(file_path, "r", encoding="utf-8", errors="ignore") as f:
        raw = f.read()

    try:
        data = json.loads(raw)
        content = json.dumps(data, indent=2)
        keys = _extract_json_keys(data)

        chunks = []
        if isinstance(data, list):
            for i, item in enumerate(data[:50]):
                chunks.append(
                    {
                        "content": json.dumps(item, indent=2),
                        "metadata": {"item_index": i, "inheritance": "array_item"},
                    }
                )
        elif isinstance(data, dict):
            for key, value in list(data.items())[:30]:
                chunks.append(
                    {
                        "content": f"{key}: {json.dumps(value, indent=2)}",
                        "metadata": {"key": key, "inheritance": f"object:{key}"},
                    }
                )

        if not chunks:
            chunks = _chunk_text(content, source=filename)

        return {
            "content": content,
            "metadata": {"type": "json", "filename": filename, "keys": keys[:20]},
            "chunks": chunks,
            "file_type": "json",
        }
    except json.JSONDecodeError:
        return _extract_text_path(file_path, filename)


def _extract_csv_path(file_path: str, filename: str) -> Dict[str, Any]:
    """Extract CSV file with row awareness."""
    import csv

    rows = []
    with open(file_path, "r", encoding="utf-8", errors="ignore") as f:
        reader = csv.reader(f)
        headers = next(reader, [])
        for row in reader:
            rows.append(row)

    content = f"Headers: {', '.join(headers)}\nRows: {len(rows)}\n\n"
    for i, row in enumerate(rows[:100]):
        content += f"Row {i + 1}: {', '.join(row)}\n"

    chunks = []
    for i, row in enumerate(rows[:200]):
        row_dict = {h: v for h, v in zip(headers, row, strict=False) if v}
        chunks.append(
            {
                "content": json.dumps(row_dict),
                "metadata": {"row_index": i, "inheritance": "csv_row"},
            }
        )

    return {
        "content": content,
        "metadata": {"type": "csv", "filename": filename, "headers": headers, "row_count": len(rows)},
        "chunks": chunks if chunks else _chunk_text(content, source=filename),
        "file_type": "csv",
    }


def _extract_pdf_path_with_ocr(file_path: str, filename: str) -> Dict[str, Any]:
    """
    Extract PDF with full support:
    - Text extraction (pdfplumber)
    - Table extraction
    - OCR fallback for image-only/scanned pages
    """
    try:
        import pdfplumber
    except ImportError:
        return _extract_pdf_path_basic(file_path, filename)

    pages_data = []
    full_text = ""
    total_tables = 0
    ocr_used = False

    try:
        with pdfplumber.open(file_path) as pdf:
            for i, page in enumerate(pdf.pages):
                page_num = i + 1
                page_text = page.extract_text() or ""

                tables = page.extract_tables() or []
                table_text = ""
                for t_idx, table in enumerate(tables):
                    total_tables += 1
                    table_text += f"\n[Table {t_idx + 1}]\n"
                    for row in table:
                        table_text += " | ".join(str(cell or "") for cell in row) + "\n"

                combined = page_text + table_text

                if len(combined.strip()) < 50:
                    ocr_text = _extract_page_with_ocr(file_path, page_num)
                    if ocr_text:
                        combined = ocr_text
                        ocr_used = True

                pages_data.append(
                    {
                        "page": page_num,
                        "content": combined,
                        "has_tables": len(tables) > 0,
                        "table_count": len(tables),
                    }
                )
                full_text += f"\n--- Page {page_num} ---\n{combined}"

        chunks = []
        for page in pages_data:
            if page["content"].strip():
                chunks.append(
                    {
                        "content": page["content"].strip(),
                        "metadata": {
                            "page": page["page"],
                            "has_tables": page["has_tables"],
                            "inheritance": f"page:{page['page']}",
                        },
                    }
                )

        return {
            "content": full_text.strip(),
            "metadata": {
                "type": "pdf",
                "filename": filename,
                "pages": len(pages_data),
                "tables": total_tables,
                "ocr_used": ocr_used,
            },
            "chunks": chunks if chunks else _chunk_text(full_text, source=filename),
            "file_type": "pdf",
        }

    except Exception as e:
        logger.error(f"PDF extraction failed: {e}")
        return _extract_pdf_path_basic(file_path, filename)


def _extract_page_with_ocr(pdf_path: str, page_num: int) -> str:
    """Extract text from a PDF page using OCR."""
    try:
        import pytesseract
        from pdf2image import convert_from_path
    except ImportError:
        return ""

    try:
        images = convert_from_path(pdf_path, first_page=page_num, last_page=page_num, dpi=300)
        if not images:
            return ""
        text = pytesseract.image_to_string(images[0], lang="eng")
        logger.info(f"OCR extracted {len(text)} chars from page {page_num}")
        return text
    except Exception as e:
        logger.warning(f"OCR failed for page {page_num}: {e}")
        return ""


def _extract_pdf_path_basic(file_path: str, filename: str) -> Dict[str, Any]:
    """Basic PDF extraction using pypdf."""
    try:
        try:
            from pypdf import PdfReader
        except ImportError:
            from PyPDF2 import PdfReader

        reader = PdfReader(file_path)
        pages = []
        full_text = ""

        for i, page in enumerate(reader.pages):
            text = page.extract_text() or ""
            pages.append({"page": i + 1, "content": text})
            full_text += f"\n--- Page {i + 1} ---\n{text}"

        chunks = []
        for page in pages:
            if page["content"].strip():
                chunks.append(
                    {
                        "content": page["content"].strip(),
                        "metadata": {"page": page["page"], "inheritance": f"page:{page['page']}"},
                    }
                )

        return {
            "content": full_text.strip(),
            "metadata": {"type": "pdf", "filename": filename, "pages": len(pages)},
            "chunks": chunks if chunks else _chunk_text(full_text, source=filename),
            "file_type": "pdf",
        }
    except Exception as e:
        return {
            "content": f"[PDF Error: {str(e)}]",
            "metadata": {"type": "pdf", "filename": filename, "error": str(e)},
            "chunks": [],
            "file_type": "pdf",
        }


def _extract_docx_path(file_path: str, filename: str) -> Dict[str, Any]:
    """Extract DOCX with structure (paragraphs, tables, headers)."""
    try:
        from docx import Document
    except ImportError:
        return {
            "content": "[DOCX Support Requires: pip install python-docx]",
            "metadata": {"type": "docx", "filename": filename, "error": "python-docx not installed"},
            "chunks": [],
            "file_type": "docx",
        }

    try:
        doc = Document(file_path)
        paragraphs = []
        for para in doc.paragraphs:
            if para.text.strip():
                style = para.style.name if para.style else "Normal"
                paragraphs.append({"text": para.text, "style": style, "is_heading": "Heading" in style})

        tables_text = []
        for t_idx, table in enumerate(doc.tables):
            table_content = f"\n[Table {t_idx + 1}]\n"
            for row in table.rows:
                row_text = " | ".join(cell.text.strip() for cell in row.cells)
                table_content += row_text + "\n"
            tables_text.append(table_content)

        content = "\n\n".join(p["text"] for p in paragraphs)
        content += "\n" + "\n".join(tables_text)

        chunks = []
        current_section = None
        section_content = []

        for p in paragraphs:
            if p["is_heading"]:
                if section_content:
                    chunks.append(
                        {
                            "content": "\n".join(section_content),
                            "metadata": {
                                "section": current_section or "Introduction",
                                "inheritance": f"section:{current_section or 'intro'}",
                            },
                        }
                    )
                current_section = p["text"]
                section_content = [p["text"]]
            else:
                section_content.append(p["text"])

        if section_content:
            chunks.append(
                {
                    "content": "\n".join(section_content),
                    "metadata": {
                        "section": current_section or "Content",
                        "inheritance": f"section:{current_section or 'content'}",
                    },
                }
            )

        for t_idx, table in enumerate(tables_text):
            chunks.append(
                {
                    "content": table,
                    "metadata": {"table_index": t_idx, "inheritance": f"table:{t_idx}"},
                }
            )

        return {
            "content": content,
            "metadata": {
                "type": "docx",
                "filename": filename,
                "paragraphs": len(paragraphs),
                "tables": len(doc.tables),
            },
            "chunks": chunks if chunks else _chunk_text(content, source=filename),
            "file_type": "docx",
        }
    except Exception as e:
        logger.error(f"DOCX extraction failed: {e}")
        return {
            "content": f"[DOCX Error: {str(e)}]",
            "metadata": {"type": "docx", "filename": filename, "error": str(e)},
            "chunks": [],
            "file_type": "docx",
        }


def _extract_xlsx_path(file_path: str, filename: str) -> Dict[str, Any]:
    """Extract XLSX with structure."""
    try:
        from openpyxl import load_workbook
    except ImportError:
        return {
            "content": "[XLSX Support Requires: pip install openpyxl]",
            "metadata": {"type": "xlsx", "filename": filename, "error": "openpyxl not installed"},
            "chunks": [],
            "file_type": "xlsx",
        }

    try:
        wb = load_workbook(file_path, data_only=True)
        sheets_data = []
        full_content = ""

        for sheet_name in wb.sheetnames:
            sheet = wb[sheet_name]
            headers = []
            rows_data = []

            for row_idx, row in enumerate(sheet.iter_rows(values_only=True)):
                if row_idx == 0:
                    headers = [str(c) if c else f"Col{i}" for i, c in enumerate(row)]
                else:
                    row_dict = {h: str(v) if v else "" for h, v in zip(headers, row, strict=False)}
                    if any(row_dict.values()):
                        rows_data.append(row_dict)

            sheet_content = f"\n=== Sheet: {sheet_name} ===\n"
            sheet_content += f"Headers: {', '.join(headers)}\nRows: {len(rows_data)}\n\n"
            for i, row in enumerate(rows_data[:50]):
                sheet_content += f"Row {i + 1}: {json.dumps(row)}\n"

            sheets_data.append({"name": sheet_name, "headers": headers, "rows": rows_data, "content": sheet_content})
            full_content += sheet_content

        chunks = []
        for sheet in sheets_data:
            if len(sheet["rows"]) <= 20:
                chunks.append(
                    {
                        "content": sheet["content"],
                        "metadata": {
                            "sheet": sheet["name"],
                            "row_count": len(sheet["rows"]),
                            "inheritance": f"sheet:{sheet['name']}",
                        },
                    }
                )
            else:
                for i, row in enumerate(sheet["rows"][:100]):
                    chunks.append(
                        {
                            "content": json.dumps(row),
                            "metadata": {
                                "sheet": sheet["name"],
                                "row_index": i,
                                "inheritance": f"sheet:{sheet['name']}:row:{i}",
                            },
                        }
                    )

        return {
            "content": full_content,
            "metadata": {
                "type": "xlsx",
                "filename": filename,
                "sheets": len(sheets_data),
                "total_rows": sum(len(s["rows"]) for s in sheets_data),
            },
            "chunks": chunks if chunks else _chunk_text(full_content, source=filename),
            "file_type": "xlsx",
        }
    except Exception as e:
        logger.error(f"XLSX extraction failed: {e}")
        return {
            "content": f"[XLSX Error: {str(e)}]",
            "metadata": {"type": "xlsx", "filename": filename, "error": str(e)},
            "chunks": [],
            "file_type": "xlsx",
        }


def _extract_pptx_path(file_path: str, filename: str) -> Dict[str, Any]:
    """Extract PPTX with structure."""
    try:
        from pptx import Presentation
    except ImportError:
        return {
            "content": "[PPTX Support Requires: pip install python-pptx]",
            "metadata": {"type": "pptx", "filename": filename, "error": "python-pptx not installed"},
            "chunks": [],
            "file_type": "pptx",
        }

    try:
        prs = Presentation(file_path)
        slides_data = []
        full_content = ""

        for slide_idx, slide in enumerate(prs.slides):
            slide_num = slide_idx + 1
            texts = []
            for shape in slide.shapes:
                if hasattr(shape, "text") and shape.text.strip():
                    texts.append(shape.text.strip())
                if shape.has_table:
                    table_text = "[Table]\n"
                    for row in shape.table.rows:
                        row_text = " | ".join(cell.text.strip() for cell in row.cells)
                        table_text += row_text + "\n"
                    texts.append(table_text)

            notes = ""
            if slide.has_notes_slide and slide.notes_slide.notes_text_frame:
                notes = slide.notes_slide.notes_text_frame.text.strip()

            slide_content = f"\n=== Slide {slide_num} ===\n" + "\n".join(texts)
            if notes:
                slide_content += f"\n[Notes]: {notes}"

            slides_data.append({"slide": slide_num, "content": slide_content, "notes": notes})
            full_content += slide_content + "\n"

        chunks = []
        for slide in slides_data:
            if slide["content"].strip():
                chunks.append(
                    {
                        "content": slide["content"].strip(),
                        "metadata": {
                            "slide": slide["slide"],
                            "has_notes": bool(slide["notes"]),
                            "inheritance": f"slide:{slide['slide']}",
                        },
                    }
                )

        return {
            "content": full_content,
            "metadata": {"type": "pptx", "filename": filename, "slides": len(slides_data)},
            "chunks": chunks if chunks else _chunk_text(full_content, source=filename),
            "file_type": "pptx",
        }
    except Exception as e:
        logger.error(f"PPTX extraction failed: {e}")
        return {
            "content": f"[PPTX Error: {str(e)}]",
            "metadata": {"type": "pptx", "filename": filename, "error": str(e)},
            "chunks": [],
            "file_type": "pptx",
        }


def _extract_doc_legacy_path(file_path: str, filename: str) -> Dict[str, Any]:
    """Extract legacy .doc files using antiword/catdoc/textract."""
    content = ""
    method = "unknown"

    try:
        result = subprocess.run(["antiword", file_path], capture_output=True, text=True, timeout=30)
        if result.returncode == 0:
            content = result.stdout
            method = "antiword"
    except (subprocess.TimeoutExpired, FileNotFoundError):
        pass

    if not content:
        try:
            result = subprocess.run(["catdoc", file_path], capture_output=True, text=True, timeout=30)
            if result.returncode == 0:
                content = result.stdout
                method = "catdoc"
        except (subprocess.TimeoutExpired, FileNotFoundError):
            pass

    if not content:
        try:
            import textract

            content = textract.process(file_path).decode("utf-8", errors="ignore")
            method = "textract"
        except ImportError:
            pass
        except Exception as e:
            logger.warning(f"textract failed: {e}")

    if not content:
        return {
            "content": "[DOC extraction requires: sudo apt-get install antiword catdoc OR pip install textract]",
            "metadata": {"type": "doc", "filename": filename, "error": "No .doc parser available"},
            "chunks": [],
            "file_type": "doc",
        }

    return {
        "content": content,
        "metadata": {"type": "doc", "filename": filename, "method": method},
        "chunks": _chunk_text(content, source=filename),
        "file_type": "doc",
    }


def _extract_xls_legacy_path(file_path: str, filename: str) -> Dict[str, Any]:
    """Extract legacy .xls files using xlrd."""
    try:
        import xlrd
    except ImportError:
        return {
            "content": "[XLS Support Requires: pip install xlrd]",
            "metadata": {"type": "xls", "filename": filename, "error": "xlrd not installed"},
            "chunks": [],
            "file_type": "xls",
        }

    try:
        wb = xlrd.open_workbook(file_path)
        sheets_data = []
        full_content = ""

        for sheet_name in wb.sheet_names():
            sheet = wb.sheet_by_name(sheet_name)
            headers = []
            rows_data = []

            for row_idx in range(sheet.nrows):
                row = [str(sheet.cell_value(row_idx, col_idx)) for col_idx in range(sheet.ncols)]
                if row_idx == 0:
                    headers = row
                else:
                    row_dict = {h: v for h, v in zip(headers, row, strict=False) if v}
                    if any(row_dict.values()):
                        rows_data.append(row_dict)

            sheet_content = f"\n=== Sheet: {sheet_name} ===\n"
            sheet_content += f"Headers: {', '.join(headers)}\nRows: {len(rows_data)}\n\n"
            for i, row in enumerate(rows_data[:50]):
                sheet_content += f"Row {i + 1}: {json.dumps(row)}\n"

            sheets_data.append({"name": sheet_name, "headers": headers, "rows": rows_data, "content": sheet_content})
            full_content += sheet_content

        chunks = []
        for sheet in sheets_data:
            if len(sheet["rows"]) <= 20:
                chunks.append(
                    {
                        "content": sheet["content"],
                        "metadata": {
                            "sheet": sheet["name"],
                            "row_count": len(sheet["rows"]),
                            "inheritance": f"sheet:{sheet['name']}",
                        },
                    }
                )
            else:
                for i, row in enumerate(sheet["rows"][:100]):
                    chunks.append(
                        {
                            "content": json.dumps(row),
                            "metadata": {
                                "sheet": sheet["name"],
                                "row_index": i,
                                "inheritance": f"sheet:{sheet['name']}:row:{i}",
                            },
                        }
                    )

        return {
            "content": full_content,
            "metadata": {
                "type": "xls",
                "filename": filename,
                "sheets": len(sheets_data),
                "total_rows": sum(len(s["rows"]) for s in sheets_data),
            },
            "chunks": chunks if chunks else _chunk_text(full_content, source=filename),
            "file_type": "xls",
        }
    except Exception as e:
        logger.error(f"XLS extraction failed: {e}")
        return {
            "content": f"[XLS Error: {str(e)}]",
            "metadata": {"type": "xls", "filename": filename, "error": str(e)},
            "chunks": [],
            "file_type": "xls",
        }


def _extract_ppt_legacy_path(file_path: str, filename: str) -> Dict[str, Any]:
    """Extract legacy .ppt files using catppt/textract."""
    content = ""
    method = "unknown"

    try:
        result = subprocess.run(["catppt", file_path], capture_output=True, text=True, timeout=30)
        if result.returncode == 0:
            content = result.stdout
            method = "catppt"
    except (subprocess.TimeoutExpired, FileNotFoundError):
        pass

    if not content:
        try:
            import textract

            content = textract.process(file_path).decode("utf-8", errors="ignore")
            method = "textract"
        except ImportError:
            pass
        except Exception as e:
            logger.warning(f"textract failed for PPT: {e}")

    if not content:
        return {
            "content": "[PPT extraction requires: sudo apt-get install catdoc OR pip install textract]",
            "metadata": {"type": "ppt", "filename": filename, "error": "No .ppt parser available"},
            "chunks": [],
            "file_type": "ppt",
        }

    slides = []
    current_slide = {"num": 1, "content": ""}
    for line in content.split("\n"):
        if line.startswith("Slide ") and line.strip().endswith(":"):
            if current_slide["content"].strip():
                slides.append(current_slide)
            try:
                slide_num = int(line.replace("Slide ", "").replace(":", "").strip())
                current_slide = {"num": slide_num, "content": ""}
            except ValueError:
                current_slide["content"] += line + "\n"
        else:
            current_slide["content"] += line + "\n"
    if current_slide["content"].strip():
        slides.append(current_slide)

    chunks = []
    for slide in slides:
        if slide["content"].strip():
            chunks.append(
                {
                    "content": slide["content"].strip(),
                    "metadata": {"slide": slide["num"], "inheritance": f"slide:{slide['num']}"},
                }
            )

    return {
        "content": content,
        "metadata": {"type": "ppt", "filename": filename, "method": method, "slides": len(slides)},
        "chunks": chunks if chunks else _chunk_text(content, source=filename),
        "file_type": "ppt",
    }


def _extract_code_path(file_path: str, filename: str) -> Dict[str, Any]:
    """Extract code file with language awareness."""
    ext = Path(filename).suffix.lower()
    language = LANGUAGE_MAP.get(ext, "text")

    with open(file_path, "r", encoding="utf-8", errors="ignore") as f:
        content = f.read()

    chunks = _chunk_code(content, language, filename)

    return {
        "content": content,
        "metadata": {"type": "code", "language": language, "filename": filename, "line_count": content.count("\n") + 1},
        "chunks": chunks,
        "file_type": f"code:{language}",
    }


# =============================================================================
# CHUNKING UTILITIES (FAIM-OPTIMIZED)
# =============================================================================


def _chunk_text(text: str, max_chars: int = 2000, overlap: int = 200, source: str = "") -> List[Dict[str, Any]]:
    """Split text into overlapping chunks optimized for FAIM."""
    if not text:
        return []

    chunks = []
    start = 0
    chunk_idx = 0

    while start < len(text):
        end = start + max_chars

        if end < len(text):
            for sep in [". ", ".\n", "\n\n", "\n", " "]:
                pos = text.rfind(sep, start, end)
                if pos > start + max_chars // 2:
                    end = pos + len(sep)
                    break

        chunk = text[start:end].strip()
        if chunk:
            chunks.append(
                {
                    "content": chunk,
                    "metadata": {
                        "chunk_index": chunk_idx,
                        "start_char": start,
                        "end_char": end,
                        "inheritance": f"chunk:{chunk_idx}",
                        "source": source,
                    },
                }
            )
            chunk_idx += 1

        start = end - overlap
        if start >= len(text):
            break

    return chunks


def _chunk_code(content: str, language: str, filename: str) -> List[Dict[str, Any]]:
    """Split code into logical chunks (functions, classes)."""
    lines = content.split("\n")
    chunks = []
    current_chunk = []
    current_def = None
    chunk_idx = 0

    def_patterns = {
        "python": ["def ", "class ", "async def "],
        "javascript": ["function ", "const ", "class ", "async function "],
        "typescript": ["function ", "const ", "class ", "interface ", "type ", "async function "],
        "go": ["func ", "type "],
        "rust": ["fn ", "struct ", "enum ", "impl ", "trait "],
        "java": ["public ", "private ", "protected ", "class ", "interface "],
    }

    patterns = def_patterns.get(language, [])

    for line in lines:
        stripped = line.strip()
        is_new_def = any(stripped.startswith(p) for p in patterns)

        if is_new_def and current_chunk:
            chunks.append(
                {
                    "content": "\n".join(current_chunk),
                    "metadata": {
                        "chunk_index": chunk_idx,
                        "definition": current_def,
                        "language": language,
                        "inheritance": f"code:{language}:{current_def or chunk_idx}",
                    },
                }
            )
            chunk_idx += 1
            current_chunk = [line]
            current_def = stripped.split("(")[0].split(":")[0] if "(" in stripped or ":" in stripped else stripped[:50]
        else:
            current_chunk.append(line)

            if len("\n".join(current_chunk)) > 1500:
                chunks.append(
                    {
                        "content": "\n".join(current_chunk),
                        "metadata": {
                            "chunk_index": chunk_idx,
                            "definition": current_def,
                            "language": language,
                            "inheritance": f"code:{language}:{current_def or chunk_idx}",
                        },
                    }
                )
                chunk_idx += 1
                current_chunk = []
                current_def = None

    if current_chunk:
        chunks.append(
            {
                "content": "\n".join(current_chunk),
                "metadata": {
                    "chunk_index": chunk_idx,
                    "definition": current_def,
                    "language": language,
                    "inheritance": f"code:{language}:{current_def or chunk_idx}",
                },
            }
        )

    return chunks if chunks else _chunk_text(content, source=filename)


def _extract_json_keys(obj, prefix: str = "") -> List[str]:
    """Recursively extract JSON keys."""
    keys = []
    if isinstance(obj, dict):
        for k, v in obj.items():
            path = f"{prefix}.{k}" if prefix else k
            keys.append(path)
            keys.extend(_extract_json_keys(v, path))
    elif isinstance(obj, list) and obj:
        keys.extend(_extract_json_keys(obj[0], f"{prefix}[]"))
    return keys[:50]
